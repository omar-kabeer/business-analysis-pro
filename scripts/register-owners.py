#!/usr/bin/env python3
"""Regenerate the Skill Reference column (L) of BABOK_v3_Artefacts_Register.xlsx.

Each skill-able row gets its owning OS skill first, from the artefact families in
docs/babok-coverage.md and the decisions in docs/skill-bindings.md; external plugin
references follow as "bind: <plugin>". Dry run by default; pass --write to save.
Requires openpyxl. Run from the repository root.
"""
import openpyxl,re,os,collections,json,sys
P='BABOK_v3_Artefacts_Register.xlsx'
wb=openpyxl.load_workbook(P); ws=wb['BABOK Artefacts Register']
skills=set(os.listdir('skills'))
EXT={'figma:figma-generate-diagram':'visual-modelling','figma:figma-design-to-code':'visual-modelling','frontend-design':'prototyping','design:user-research':'ux','data:analyze':'data-analysis','data:write-query':'business-intelligence','data:build-dashboard':'business-intelligence','data:data-context-extractor':'business-intelligence','data:statistical-analysis':'data-analysis','engineering:architecture':'architecture','engineering:code-review':'quality','engineering:testing-strategy':'acceptance-testing','legal:review-contract':'procurement-contracts','legal:compliance-check':'regulatory-compliance','operations:risk-assessment':'risk-analysis','operations:compliance-tracking':'regulatory-compliance','operations:change-request':'change-control','operations:capacity-plan':'estimation','operations:process-optimization':'process-modelling','operations:vendor-review':'vendor-evaluation','operations:runbook':'technical-writer','enterprise-search:search':'document-analysis','productivity:task-management':'governance','product-management:roadmap-update':'product-manager','finance:variance-analysis':'finance','finance:financial-statements':'finance'}
# ordered family rules from docs/babok-coverage.md (artefact name, lowercase)
RULES=[
 (r'business analysis approach|stakeholder engagement approach|governance approach|business analysis plan|business analysis performance','ba-planning'),
 (r'information management approach|artefact register|traceability repository|information architecture|traceability repository|requirements repository','information-management'),
 (r'traceab|requirements architecture','requirements'),
 (r'interview|workshop|focus[- ]group|survey|questionnaire|observation|brainstorm|collaborative game|elicitation|stakeholder list|stakeholder map|stakeholder analysis|stakeholder impact','elicitation'),
 (r'existing document|document analysis|domain knowledge|supporting material','document-analysis'),
 (r'capability map|capability model|value stream|value chain|organi[sz]ational (map|model|chart)|org chart|information map|reference model|framework|zachman|togaf|archimate|business architecture|roadmap \(enterprise\)','business-architecture'),
 (r'bpmn|swimlane|sipoc|process (model|architecture|performance|analysis)|flowchart|activity diagram|value stream map|functional decomposition|decomposition','process-modelling'),
 (r'entity|class (model|diagram)|data dictionary|concept model|data flow|data model|metadata','data-modelling'),
 (r'benchmark|market analysis|competitive|market research','market-research'),
 (r'warehouse|data mart|etl|dashboard|scorecard|reports? (and|&) charts|reporting|data quality|data sources|data mining','business-intelligence'),
 (r'solution performance|limitation|recommended action|lessons learned','solution-evaluation'),
 (r'metric|kpi|performance measure|performance analysis','data-analysis'),
 (r'wireframe|mock-up|storyboard|prototype|proof of concept','prototyping'),
 (r'persona|journey|empathy|usability|accessib','ux'),
 (r'design option|interface|solution architecture|technical design|blueprint|sequence diagram|state (model|diagram|table)|service-oriented','architecture'),
 (r'decision (matrix|table|tree|model|analysis)|expected value','decision-analysis'),
 (r'estimat|work breakdown','estimation'),
 (r'business case|financial|cost-benefit|roi|npv','finance'),
 (r'request for information|request for proposal|rfi|rfp|vendor','vendor-evaluation'),
 (r'statement of work|service level|request for quote|request for tender|contract','procurement-contracts'),
 (r'inspection|walkthrough|peer review|checklist|review record|defect','quality'),
 (r'user acceptance|test','acceptance-testing'),
 (r'risk','risk-analysis'),
 (r'change (request|assessment)|impact analysis|proposed change','change-control'),
 (r'policy|policies|legal|regulat|compliance|audit','regulatory-compliance'),
 (r'raid|decision log|item (log|tracking)|issue log|release readiness','governance'),
 (r'backlog|epic|feature|user stor|release plan|definition of (ready|done)|story map|spike','product-owner'),
 (r'vision|product roadmap|discovery|prioriti|go to market|product brief','product-manager'),
 (r'safe|lightweight documentation|ceremon','agile-coach'),
 (r'status report|meeting notes|communicat|announcement','communication'),
 (r'executive summary|board paper','executive-review'),
 (r'current state|future state|gap analysis|business rule|glossary','business-analysis'),
]
def owner(art,old_os):
  a=art.lower()
  for rx,s in RULES:
    if re.search(rx,a): return s,'family'
  return (old_os[0],'kept') if old_os else (None,'none')
changes=collections.Counter(); unresolved=[]
for r in range(2,ws.max_row+1):
  art=ws.cell(r,3).value
  if not art: continue
  L=str(ws.cell(r,12).value or '')
  parts=[p.strip() for p in re.split(r';\s*',L) if p.strip()]
  old_os=[p.split(':')[1] for p in parts if p.startswith('business-analysis-os:') and p.split(':')[1] in skills]
  ext=[p for p in parts if not p.startswith('business-analysis-os:')]
  skillable=ws.cell(r,11).value
  if skillable in ('No','n/a'):
    changes['not skill-able, unchanged']+=1; continue
  own,how=('reference-standards','binding') if '(reference standard' in L else owner(art,old_os)
  if own is None:
    for e in ext:
      k=e.split(' ')[0]
      if k in EXT: own,how=EXT[k],'external'; break
  if own is None:
    unresolved.append((r,art,L)); continue
  assert own in skills,own
  new='business-analysis-os:'+own
  binds=[e for e in ext if e.split(' ')[0] in EXT]
  notes=[e for e in ext if e not in binds]
  val='; '.join([new]+[f'bind: {b}' for b in binds]+notes)
  if val!=L: changes[how]+=1
  ws.cell(r,12).value=val
print(dict(changes)); print('unresolved',unresolved)
if '--write' in sys.argv: wb.save(P)
cnt=collections.Counter()
for r in range(2,ws.max_row+1):
  v=str(ws.cell(r,12).value or '')
  if v.startswith('business-analysis-os:'): cnt[v.split(';')[0].split(':')[1]]+=1
print(len(cnt), sorted(set(skills)-set(cnt)))
